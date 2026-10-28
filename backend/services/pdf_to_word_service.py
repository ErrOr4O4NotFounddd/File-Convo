import os
import io
import shutil
import tempfile
import pymupdf as fitz
from pathlib import Path
from typing import Dict, Any, Optional
from backend.config import load_settings, DEFAULT_OUTPUT_DIR, get_output_dir
from backend.services.history_service import add_history_entry

def get_unique_output_path(output_dir: Path, filename: str) -> Path:
    stem = Path(filename).stem
    ext = Path(filename).suffix
    candidate = output_dir / f"{stem}{ext}"
    counter = 1
    while candidate.exists():
        candidate = output_dir / f"{stem}_{counter}{ext}"
        counter += 1
    return candidate

def convert_pdf_to_docx(
    file_bytes: bytes,
    original_filename: str,
    custom_output_dir: Optional[str] = None
) -> Dict[str, Any]:
    output_dir = get_output_dir(custom_output_dir)

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        page_count = len(doc)
        if page_count == 0:
            doc.close()
            err = "PDF document is empty (0 pages)."
            add_history_entry(original_filename, "PDF → DOCX", "error", error_message=err)
            return {"success": False, "error": err}

        total_text_chars = 0
        for page in doc:
            total_text_chars += len(page.get_text().strip())
        doc.close()

        is_scanned_notice = None
        if total_text_chars < 20 * page_count:
            is_scanned_notice = (
                "Notice: This PDF contains little or no selectable digital text. It appears to be scanned or image-based. "
                "For scanned PDFs, please use the 'OCR and Text Extraction' tool to extract editable text."
            )
    except Exception as e:
        err = f"Failed to parse PDF document: {str(e)}"
        add_history_entry(original_filename, "PDF → DOCX", "error", error_message=err)
        return {"success": False, "error": err}

    stem = Path(original_filename).stem
    out_filename = f"{stem}.docx"
    final_out_path = get_unique_output_path(output_dir, out_filename)

    with tempfile.TemporaryDirectory() as temp_dir_str:
        temp_dir = Path(temp_dir_str)
        temp_pdf = temp_dir / "input.pdf"
        temp_docx = temp_dir / "output.docx"
        temp_pdf.write_bytes(file_bytes)

        conversion_succeeded = False
        primary_error = None

        try:
            from pdf2docx import Converter
            cv = Converter(str(temp_pdf))
            cv.convert(str(temp_docx), start=0, end=None)
            cv.close()
            if temp_docx.exists() and temp_docx.stat().st_size > 0:
                conversion_succeeded = True
        except Exception as e:
            primary_error = str(e)

        if not conversion_succeeded:
            try:
                from docx import Document
                from docx.shared import Inches, Pt
                doc_writer = Document()
                pdf_doc = fitz.open(str(temp_pdf))

                for pno, page in enumerate(pdf_doc):
                    if pno > 0:
                        doc_writer.add_page_break()

                    text_blocks = page.get_text("blocks")
                    for b in text_blocks:
                        text = b[4].strip()
                        if text:
                            p = doc_writer.add_paragraph(text)
                            p.paragraph_format.space_after = Pt(6)

                    images = page.get_images()
                    for img_meta in images:
                        try:
                            xref = img_meta[0]
                            base_img = pdf_doc.extract_image(xref)
                            img_bytes = base_img["image"]
                            img_stream = io.BytesIO(img_bytes)
                            doc_writer.add_picture(img_stream, width=Inches(5.0))
                        except Exception:
                            pass

                pdf_doc.close()
                doc_writer.save(str(temp_docx))
                if temp_docx.exists() and temp_docx.stat().st_size > 0:
                    conversion_succeeded = True
            except Exception as fallback_err:
                err = f"PDF to Word conversion failed: {primary_error or str(fallback_err)}"
                add_history_entry(original_filename, "PDF → DOCX", "error", error_message=err)
                return {"success": False, "error": err}

        if not conversion_succeeded or not temp_docx.exists():
            err = f"Failed to generate valid DOCX: {primary_error or 'Unknown error'}"
            add_history_entry(original_filename, "PDF → DOCX", "error", error_message=err)
            return {"success": False, "error": err}

        try:
            from docx import Document
            _ = Document(str(temp_docx))
        except Exception as v_err:
            err = f"Output DOCX validation failed: {str(v_err)}"
            add_history_entry(original_filename, "PDF → DOCX", "error", error_message=err)
            return {"success": False, "error": err}

        shutil.copy2(str(temp_docx), str(final_out_path))

    out_size = final_out_path.stat().st_size
    stats = {
        "original_size": len(file_bytes),
        "output_size": out_size,
        "page_count": page_count,
        "output_format": "DOCX",
        "fidelity_notice": "Note: Exact formatting (custom fonts, intricate columns, complex tables) may differ slightly from the original PDF."
    }
    if is_scanned_notice:
        stats["scanned_warning"] = is_scanned_notice

    add_history_entry(
        original_filename=original_filename,
        conversion_type="PDF → DOCX",
        status="success",
        output_location=str(final_out_path),
        details=stats
    )

    return {
        "success": True,
        "output_path": str(final_out_path),
        "output_filename": final_out_path.name,
        "stats": stats
    }
